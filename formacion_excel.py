"""Formación y empleo de la Delegación de Empleo del Ayuntamiento de Marbella.

No hay fuente abierta: los datos son el registro que lleva la propia Delegación
(Excel del indicador SEGITTUR GOB02_05_01, «Plan de formación en turismo — sector
privado y ciudadanía»). Este script NO va en la Action: se ejecuta a mano cada vez
que la Delegación entregue el Excel actualizado y se commitea data/formacion.json.

    python formacion_excel.py ["ruta\\al\\Excel.xlsx"]

Hojas que lee:
  DATOS      — totales anuales tal como los da la Delegación (son las cifras oficiales).
  Formación  — un curso por fila; sirve para los desgloses (modalidad, área, curso).
Las dos hojas no cuadran del todo (2020 y 2023 difieren en unas decenas de alumnos):
los totales salen siempre de DATOS y los desgloses de la hoja de cursos.
"""
import datetime, io, json, os, re, sys, unicodedata
import openpyxl

EXCEL = (r"C:\Users\josem\OneDrive\Escritorio\SEGITTUR JMHF\Ejes\01_Gobernanza\B. Eficiencia en la gestión"
         r"\GOB02_05_01_R02 Plan de formacion en turismo (Sector privado y ciudadanía).xlsx")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "formacion.json")

# columnas de la hoja DATOS -> clave del JSON
COL_DATOS = {
    "Ofertas publicadas Portal Empleo": "ofertas_portal",
    "Organizaciones participantes en la Feria de Empleo": "feria_empresas",
    "Asistentes a la Feria de Empleo": "feria_asistentes",
    "Acciones formativas Presenciales o Mixtas": "acc_presencial",
    "Acciones formativas Teleformación": "acc_tele",
    "Total acciones formativas": "acciones",
    "Plazas Ofertadas": "plazas",
    "Nº solicitudes": "solicitudes",
    "Nº alumnos finalizan formación": "finalizan",
    "Horas de formación ofertadas": "horas_of",
    "Horas de formación recibidas": "horas_rec",
    "Presupuesto formación (€)": "presupuesto",
    "Importe Ejecutado (€)": "ejecutado",
}
# campos que solo existen cuando el curso ha terminado: en el año en curso un 0 es «aún no hay dato»
DE_CIERRE = ("solicitudes", "finalizan", "horas_rec", "ejecutado")

# Áreas temáticas (como mucho ocho: es el máximo de series legibles en una gráfica).
# El orden importa: gana la primera que encaje («Atención al cliente en inglés» es Idiomas).
AREAS = [
    ("Idiomas", r"ingl|franc"),
    ("Hostelería y alimentación", r"aliment|alerg|catering|carnicer|restaurac"),
    ("Comercio electrónico y marketing", r"comercio electr|publicidad|prestashop|tienda virtual|wordpress|pagina|web"),
    ("Gestión de empresa", r"contab|fiscal|factura|almacen|proteccion de datos|firma digital"),
    ("Competencias digitales y ofimática", r"ofimat|excel|microsoft|google|autocad|internet|digital|big data|videojuego|grabacion"),
    ("Jardinería y mantenimiento", r"jardin|poda|riego|fitosanit|plaga|piscina|legionel|desinfec|domot|knx|electric"),
    ("Seguridad, PRL y logística", r"riesgos|vigilante|socorr|parking|cap |cap$|carretill|plataforma"),
    ("Atención al cliente y servicios", r"atencion|coaching|empleab|peluquer|sociosanit"),
]


def plano(s):
    s = unicodedata.normalize("NFD", str(s)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


def clave_curso(nombre):
    """Agrupa las ediciones de un mismo curso aunque el nombre venga escrito distinto."""
    k = plano(nombre)
    k = re.sub(r"\(.*?\)", "", k)                      # «(190 horas)», «(MF0233_2)»
    k = re.sub(r"^(fpe(-| de )|[a-z]{4}\d{2,4}(po)?|fom\d+)\s*", "", k)   # códigos de certificado
    k = re.sub(r"-(marbella|guadalmina)$", "", k)
    k = k.replace("excell", "excel").replace("conpetencias", "competencias").replace("procolo", "protocolo")
    k = re.sub(r"[^a-z0-9]+", " ", k).strip()
    k = re.sub(r"^ms ", "", k)
    if k.startswith("manipula") and "aliment" in k:
        k = "manipulador de alimentos"
    if k.startswith("competencias digitales basicas"):
        k = "competencias digitales basicas"
    if "prestashop" in k:
        k = "prestashop"
    if k.startswith("frances elemental"):
        k = "frances elemental para pisos"
    if k.startswith("cap "):
        k = "cap"
    if "carretillas" in k:
        k = "carretillas elevadoras"
    return k


def area(nombre):
    k = plano(nombre)
    for a, rx in AREAS:
        if re.search(rx, k):
            return a
    return "Atención al cliente y servicios"


def bonito(nombre):
    """Nombre legible: sin códigos ni paréntesis y en minúsculas si venía en mayúsculas."""
    n = re.sub(r"\(.*?\)", "", str(nombre)).strip()
    n = re.sub(r"^(FPE-|FPE de |[A-Z]{4}\d{2,4}(PO)?\s+|FOM\d+\s+)", "", n)
    n = re.sub(r"-(MARBELLA|GUADALMINA)$", "", n).strip(" .-")
    n = re.sub(r"\s+", " ", n.replace("\n", " "))
    if n.isupper():
        n = n.capitalize()
    return n.replace("Excell", "Excel").replace("Conpetencias", "Competencias").replace("Procolo", "Protocolo")


def num(x):
    return x if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    hoy = datetime.date.today()

    ws = wb["DATOS"]
    cab = [c.value for c in ws[1]]
    anual = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not isinstance(row[0], int):
            continue
        r = {"y": row[0]}
        for i, h in enumerate(cab):
            if h in COL_DATOS:
                r[COL_DATOS[h]] = num(row[i])
        if r["y"] >= hoy.year:
            for k in DE_CIERRE:
                if not r.get(k):
                    r[k] = None
        r["estado"] = "cerrado" if r.get("finalizan") else ("en curso" if r["y"] >= hoy.year else "sin cierre")
        anual.append(r)

    ws = wb["Formación"]
    cursos = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not isinstance(row[0], int) or not row[1]:
            continue
        mod = str(row[2] or "").strip()
        cursos.append({
            "y": row[0], "curso": bonito(row[1]), "clave": clave_curso(row[1]), "area": area(row[1]),
            "modalidad": "Teleformación" if mod == "Teleformación" else "Presencial o mixta",
            "solicitudes": num(row[3]), "plazas": num(row[4]), "finalizan": num(row[5]),
            "horas": num(row[6]), "horas_of": num(row[7]), "horas_rec": num(row[8]),
        })
    # mismo criterio que en DATOS: en el año en curso el curso aún no tiene cierre
    for c in cursos:
        if c["y"] >= hoy.year:
            for k in ("solicitudes", "finalizan", "horas_rec"):
                if not c[k]:
                    c[k] = None

    sin_area = sorted({c["curso"] for c in cursos if not any(re.search(rx, plano(c["curso"])) for _, rx in AREAS)})
    if sin_area:
        print("  Aviso: cursos sin área propia (van a «Atención al cliente y servicios»):", *sin_area, sep="\n    ")

    out = {"generado": hoy.isoformat(), "fuente": os.path.basename(path), "anual": anual, "cursos": cursos}
    with io.open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    print(f"  data/formacion.json · {len(anual)} años ({anual[0]['y']}-{anual[-1]['y']}) · {len(cursos)} cursos · "
          f"{len({c['clave'] for c in cursos})} cursos distintos")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else EXCEL)
